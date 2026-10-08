"""本周课表：支持手动录入或从 CSV 导入课程，并打印文本视图。

CSV 格式（含表头）：课程名,星期几,开始时间,结束时间
用法：
    python schedule2.py                      # 交互菜单
    python schedule2.py courses.csv          # 直接导入 CSV 并打印
    python schedule2.py 课表.xlsx [周次]      # 直接导入 xlsx 并打印（周次缺省按今天推算）

功能：
    - 需求1：手动录入 / CSV / xlsx 导入课程，打印「时间 × 星期」网格视图
    - 需求2：给定每日可用范围（默认 08:00-22:00），算每天空闲时段；
            同一课程同一天间隔 ≤15 分钟的连堂小节自动合并，避免碎片。
"""
import csv
import datetime
import re
import sys
import unicodedata
from collections import namedtuple

# day: 0=周一 ... 6=周日；teacher/location 仅 xlsx 导入时有值
Course = namedtuple("Course", ["name", "day", "start", "end", "teacher", "location"],
                    defaults=["", ""])

SEMESTER_START = datetime.date(2026, 8, 31)  # 第 1 周的周一，用于推算“本周是第几周”

DAYS = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]

DAY_MAP = {str(i + 1): i for i in range(7)}
for i, ch in enumerate("一二三四五六日"):
    DAY_MAP[ch] = DAY_MAP["周" + ch] = DAY_MAP["星期" + ch] = i
DAY_MAP["天"] = DAY_MAP["周天"] = DAY_MAP["星期天"] = 6

TIME_RE = re.compile(r"^(\d{1,2}):([0-5]\d)$")


def parse_day(text):
    """把 '周一' / '1' / '星期三' 等写法解析成 0-6。"""
    key = text.strip()
    if key in DAY_MAP:
        return DAY_MAP[key]
    raise ValueError(f"无法识别的星期：{text!r}（可用 周一~周日 或 1~7）")


def parse_time(text):
    """把 '8:00' / '08:00' 规范成 'HH:MM'。"""
    m = TIME_RE.match(text.strip())
    if not m or int(m.group(1)) > 23:
        raise ValueError(f"无法识别的时间：{text!r}（格式 HH:MM，如 08:00）")
    return f"{int(m.group(1)):02d}:{m.group(2)}"


def time_to_minutes(t):
    """把 'HH:MM' 转成当日分钟数，便于区间运算。"""
    h, m = t.split(":")
    return int(h) * 60 + int(m)


def minutes_to_time(m):
    """把当日分钟数转回 'HH:MM'。"""
    return f"{m // 60:02d}:{m % 60:02d}"


def parse_range(text):
    """解析每日可用范围（如 '08:00-22:00'，支持 - — ~ 至 到），空串用默认 08:00-22:00。"""
    text = text.strip() or "08:00-22:00"
    m = re.match(r"^(\d{1,2}:\d{2})\s*[-—~至到]\s*(\d{1,2}:\d{2})$", text)
    if not m:
        raise ValueError("格式应为 HH:MM-HH:MM，如 08:00-22:00")
    start, end = parse_time(m.group(1)), parse_time(m.group(2))
    if time_to_minutes(end) <= time_to_minutes(start):
        raise ValueError("结束时间必须晚于开始时间")
    return start, end


def make_course(name, day, start, end):
    course = Course(name.strip(), parse_day(day), parse_time(start), parse_time(end))
    if not course.name:
        raise ValueError("课程名不能为空")
    if course.end <= course.start:
        raise ValueError(f"结束时间必须晚于开始时间：{course.start}~{course.end}")
    return course


# ---------- 录入 ----------

def input_courses():
    """手动录入，课程名留空结束。返回 Course 列表。"""
    courses = []
    print("手动录入（课程名直接回车结束）")
    while True:
        name = input("  课程名: ").strip()
        if not name:
            break
        day = input("  星期几(周一~周日 或 1~7): ")
        start = input("  开始时间(HH:MM): ")
        end = input("  结束时间(HH:MM): ")
        try:
            courses.append(make_course(name, day, start, end))
            print("  -> 已添加")
        except ValueError as e:
            print(f"  -> 输入无效：{e}，请重填本行")
    return courses


def load_csv(path):
    """从 CSV 读入课程；坏行跳过并提示。返回 Course 列表。"""
    courses = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        for lineno, row in enumerate(csv.reader(f), 1):
            if not row or not any(cell.strip() for cell in row):
                continue
            if lineno == 1 and row[0].strip() in ("课程名", "name"):
                continue  # 跳过表头
            if len(row) < 4:
                print(f"  第 {lineno} 行字段不足，已跳过：{row}")
                continue
            try:
                courses.append(make_course(row[0], row[1], row[2], row[3]))
            except ValueError as e:
                print(f"  第 {lineno} 行无效，已跳过：{e}")
    print(f"从 {path} 导入 {len(courses)} 门课程")
    return courses


# ---------- xlsx 课表（学校教务系统导出格式） ----------

# 单元格内每个课程块的格式：
#   课程名 课程号
#   1-4周,6-18周 教师 HH:MM-HH:MM 【地点】
ENTRY_RE = re.compile(
    r"^(?P<weeks>(?:\d+(?:-\d+)?周[,，]?)+)\s+"
    r"(?P<teacher>.*?)\s+"
    r"(?P<start>\d{1,2}:\d{2})-(?P<end>\d{1,2}:\d{2})"
    r"(?:\s*【(?P<location>.*)】)?\s*$")
HEADER_RE = re.compile(r"^(?P<name>.+?)\s+\d{2}$")


def current_week(today=None):
    """按 SEMESTER_START 推算今天是本学期第几周。"""
    today = today or datetime.date.today()
    return (today - SEMESTER_START).days // 7 + 1


def parse_weeks(spec):
    """'1-4周,6-18周' -> {1,2,3,4,6,7,...,18}。"""
    weeks = set()
    for part in re.split(r"[,，]", spec):
        part = part.strip()
        if not part:
            continue
        m = re.fullmatch(r"(\d+)(?:-(\d+))?周", part)
        if not m:
            raise ValueError(f"无法识别的周次：{spec!r}")
        a, b = int(m.group(1)), int(m.group(2) or m.group(1))
        weeks.update(range(a, b + 1))
    return weeks


def parse_cell(text, day, week):
    """解析一个 xlsx 单元格，返回该周上课的 Course 列表。"""
    courses, name = [], None
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        em = ENTRY_RE.match(line)
        if em:
            if name and week in parse_weeks(em.group("weeks")):
                courses.append(Course(name, day, parse_time(em.group("start")),
                                      parse_time(em.group("end")),
                                      em.group("teacher"), em.group("location") or ""))
            continue
        hm = HEADER_RE.match(line)
        if hm:
            name = hm.group("name").strip()
    return courses


def load_xlsx(path, week):
    """从教务系统导出的 xlsx 课表读入指定周次的课程。返回 Course 列表。"""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    courses = []
    for ws in wb.worksheets:
        header_row, col_day = None, {}
        for row in ws.iter_rows():
            days = {c.column: DAY_MAP[str(c.value).strip()] for c in row
                    if isinstance(c.value, str)
                    and str(c.value).strip().startswith("星期")
                    and str(c.value).strip() in DAY_MAP}
            if len(days) >= 7:
                header_row, col_day = row[0].row, days
                break
        if header_row is None:
            continue
        for row in ws.iter_rows(min_row=header_row + 1):
            for c in row:
                if c.column in col_day and isinstance(c.value, str) and "周" in c.value:
                    courses.extend(parse_cell(c.value, col_day[c.column], week))
    courses.sort(key=lambda c: (c.day, c.start))
    print(f"从 {path} 导入第 {week} 周的 {len(courses)} 门课程")
    return courses


# ---------- 打印 ----------

def merge_same_course(courses, max_gap=15):
    """同一课程同一天、间隔不超过 max_gap 分钟的连续小节合并为一个时段。

    用于消除 09:00-09:45、09:55-10:40 这类碎片，合并后显示为 09:00-10:40。
    """
    groups = {}
    for c in courses:
        groups.setdefault((c.day, c.name), []).append(c)
    merged = []
    for _, items in groups.items():
        items.sort(key=lambda c: (c.start, c.end))
        cur = items[0]
        for nxt in items[1:]:
            if time_to_minutes(nxt.start) - time_to_minutes(cur.end) <= max_gap:
                cur = cur._replace(end=minutes_to_time(
                    max(time_to_minutes(cur.end), time_to_minutes(nxt.end))))
            else:
                merged.append(cur)
                cur = nxt
        merged.append(cur)
    return sorted(merged, key=lambda c: (c.day, c.start))


def compute_free_slots(courses, day_start, day_end):
    """在每日可用范围内扣除课程占用，返回 {day: [(start_min, end_min), ...]}。"""
    busy = {d: [] for d in range(7)}
    for c in merge_same_course(courses):
        busy[c.day].append((time_to_minutes(c.start), time_to_minutes(c.end)))
    lo, hi = time_to_minutes(day_start), time_to_minutes(day_end)
    free = {}
    for d in range(7):
        slots, cur = [], lo
        for bs, be in sorted(busy[d]):
            if be <= lo or bs >= hi:
                continue
            if max(bs, lo) > cur:
                slots.append((cur, max(bs, lo)))
            cur = max(cur, min(be, hi))
        if cur < hi:
            slots.append((cur, hi))
        free[d] = slots
    return free


def print_free_slots(courses, day_start, day_end):
    """打印每天的空闲时段。"""
    print(f"\n每日可用范围 {day_start}-{day_end} 内的空闲时段：")
    full = (time_to_minutes(day_start), time_to_minutes(day_end))
    free = compute_free_slots(courses, day_start, day_end)
    for d in range(7):
        if free[d] == [full]:
            print(f"  {DAYS[d]}: {day_start}-{day_end}（全天空闲）")
        elif not free[d]:
            print(f"  {DAYS[d]}: 无空闲")
        else:
            print(f"  {DAYS[d]}: " + ", ".join(
                f"{minutes_to_time(a)}-{minutes_to_time(b)}" for a, b in free[d]))


def display_width(s):
    """显示宽度：中文等全角字符按 2 计。"""
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in s)


def pad(s, width):
    return s + " " * max(0, width - display_width(s))


def print_schedule(courses):
    """按 时间 × 星期 的网格打印本周课表。"""
    print("\n本周课表")
    if not courses:
        print("（暂无课程）")
        return
    courses = merge_same_course(courses)
    slots = sorted({(c.start, c.end) for c in courses})
    cell = {}
    for c in courses:
        key = (c.start, c.end, c.day)
        cell[key] = cell[key] + "/" + c.name if key in cell else c.name

    headers = ["时间"] + DAYS
    widths = [max(display_width(f"{s}-{e}") for s, e in slots)]
    widths += [max([display_width(DAYS[d])] +
                   [display_width(cell.get((s, e, d), "")) for s, e in slots])
               for d in range(7)]

    def render(cells):
        return " | ".join(pad(c, w) for c, w in zip(cells, widths))

    print(render(headers))
    print("-+-".join("-" * w for w in widths))
    for s, e in slots:
        print(render([f"{s}-{e}"] + [cell.get((s, e, d), "") for d in range(7)]))

    detailed = [c for c in sorted(courses, key=lambda c: (c.day, c.start))
                if c.teacher or c.location]
    if detailed:
        print("\n课程详情：")
        for c in detailed:
            line = f"  {DAYS[c.day]} {c.start}-{c.end} {c.name}"
            if c.teacher:
                line += f"  {c.teacher}"
            if c.location:
                line += f"  【{c.location}】"
            print(line)


def export_schedule_xlsx(courses, output_path):
    """将本周课表导出为带样式的 xlsx 文件。"""
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "本周课表"

    headers = ["时间"] + DAYS
    thin = Side(border_style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True, size=12)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2",
                              fill_type="solid")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 标题行
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    if not courses:
        ws.cell(row=2, column=1, value="（暂无课程）")
    else:
        slots = sorted({(c.start, c.end) for c in courses})
        cell_map = {}
        for c in courses:
            key = (c.start, c.end, c.day)
            cell_map[key] = cell_map[key] + "/" + c.name if key in cell_map else c.name

        for r, (s, e) in enumerate(slots, 2):
            tc = ws.cell(row=r, column=1, value=f"{s}-{e}")
            tc.alignment = center
            tc.border = border
            for d in range(7):
                cc = ws.cell(row=r, column=d + 2,
                             value=cell_map.get((s, e, d), ""))
                cc.alignment = center
                cc.border = border

        # 列宽自适应（中文按 2 计）
        ws.column_dimensions["A"].width = max(display_width(f"{s}-{e}")
                                              for s, e in slots) + 2
        for d in range(7):
            col_letter = openpyxl.utils.get_column_letter(d + 2)
            max_w = display_width(DAYS[d])
            for s, e in slots:
                max_w = max(max_w, display_width(cell_map.get((s, e, d), "")))
            ws.column_dimensions[col_letter].width = max_w + 2
        ws.row_dimensions[1].height = 24

    # 课程详情 sheet
    detailed = [c for c in sorted(courses, key=lambda c: (c.day, c.start))
                if c.teacher or c.location]
    if detailed:
        ws2 = wb.create_sheet("课程详情")
        detail_headers = ["星期", "时间", "课程名", "教师", "地点"]
        for col, h in enumerate(detail_headers, 1):
            cell = ws2.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border
        for r, c in enumerate(detailed, 2):
            values = [DAYS[c.day], f"{c.start}-{c.end}", c.name,
                      c.teacher or "", c.location or ""]
            for col, v in enumerate(values, 1):
                cell = ws2.cell(row=r, column=col, value=v)
                cell.alignment = center
                cell.border = border
        for col in range(1, 6):
            ws2.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 16
        ws2.row_dimensions[1].height = 24

    wb.save(output_path)
    print(f"课表已导出到：{output_path}")


# ---------- 主程序 ----------

def main():
    if len(sys.argv) > 1:
        path = sys.argv[1]
        if path.lower().endswith(".xlsx"):
            week = int(sys.argv[2]) if len(sys.argv) > 2 else current_week()
            print(f"本学期第 1 周始于 {SEMESTER_START}，按第 {week} 周筛选")
            courses = load_xlsx(path, week)
            output = sys.argv[3] if len(sys.argv) > 3 else None
        else:
            courses = load_csv(path)
            output = sys.argv[2] if len(sys.argv) > 2 else None
        print_schedule(courses)
        if output:
            export_schedule_xlsx(courses, output)
        return

    courses = []
    while True:
        print("\n===== 课表 =====")
        print("1. 手动录入课程")
        print("2. 从 CSV 导入")
        print("3. 从 xlsx 课表导入（按周次筛选）")
        print("4. 打印本周课表")
        print("5. 导出本周课表为 Excel")
        print("6. 查询每天的空闲时段")
        print("0. 退出")
        choice = input("请选择: ").strip()
        if choice == "1":
            courses.extend(input_courses())
        elif choice == "2":
            path = input("CSV 文件路径: ").strip().strip('"')
            try:
                courses.extend(load_csv(path))
            except OSError as e:
                print(f"读取失败：{e}")
        elif choice == "3":
            path = input("xlsx 文件路径: ").strip().strip('"')
            week_s = input(f"第几周(直接回车按今天推算: 第{current_week()}周): ").strip()
            try:
                week = int(week_s) if week_s else current_week()
                courses.extend(load_xlsx(path, week))
            except (OSError, ValueError) as e:
                print(f"读取失败：{e}")
        elif choice == "4":
            print_schedule(courses)
        elif choice == "5":
            if not courses:
                print("请先导入或录入课程")
                continue
            out = input("输出文件名(默认 本周课表.xlsx): ").strip().strip('"') or "本周课表.xlsx"
            export_schedule_xlsx(courses, out)
        elif choice == "6":
            if not courses:
                print("请先导入或录入课程")
                continue
            range_s = input("每日可用范围(如 08:00-22:00，回车默认): ").strip()
            try:
                day_start, day_end = parse_range(range_s)
                print_free_slots(courses, day_start, day_end)
            except ValueError as e:
                print(f"输入无效：{e}")
        elif choice == "0":
            break
        else:
            print("无效选择")


if __name__ == "__main__":
    main()
